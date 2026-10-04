import base64
import io
import zipfile
from pathlib import Path
from fastapi import HTTPException
from pydantic import ValidationError
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from .schemas import Transcript

MEDIA = {'.mp3':'audio/mpeg','.wav':'audio/wav','.m4a':'audio/mp4',
         '.ogg':'audio/ogg','.flac':'audio/flac','.mp4':'video/mp4','.webm':'video/webm'}
SUPPORTED = {'.txt','.md','.csv','.pdf','.docx','.pptx',*MEDIA}
MAX_FILE = 8 * 1024 * 1024


def check_zip(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if len(z.infolist()) > 3000 or sum(x.file_size for x in z.infolist()) > 60*1024*1024:
            raise HTTPException(413, 'Office file expands beyond the demo limit.')

async def extract(name, data, provider):
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise HTTPException(415, f'Unsupported file type: {ext}. Save older Word/PowerPoint files as DOCX/PPTX.')
    if not data or len(data) > MAX_FILE:
        raise HTTPException(413, 'Each file must be nonempty and at most 8 MB.')
    warnings, segments = [], []
    try:
        if ext in MEDIA:
            value = await provider.json('''Transcribe the supplied media in its original language.
Return {"segments":[{"start":0.0,"end":2.5,"text":"exact spoken words"}],
"visual_description":""}. Times are seconds from the beginning, chronological,
non-overlapping, with end > start. Do not fabricate speech. For video, also give
an objective visual description; do not infer hidden events or identities.
For audio leave visual_description empty. Keep the transcription complete.
If there is no intelligible speech return an empty segments list.''',
                media={'mimeType':MEDIA[ext], 'data':base64.b64encode(data).decode()})
            transcript = Transcript.model_validate(value)
            last_end = 0.0
            for s in transcript.segments:
                if s.end <= s.start or s.start < last_end:
                    raise HTTPException(502, 'Invalid/overlapping AI timestamps. Retry a shorter clip.')
                last_end = s.end
                segments.append(s.model_dump())
            chunks = [(f'audio {s.start:.2f}–{s.end:.2f}s', s.text) for s in transcript.segments]
            if transcript.visual_description:
                chunks.append(('AI visual description', transcript.visual_description))
            warnings.append('AI transcription/visual description and timestamps are approximate. Check against the original media. Use short clips (about 2 minutes) for this demo.')
        elif ext in {'.txt','.md','.csv'}:
            chunks = [('text', data.decode('utf-8-sig'))]
        elif ext == '.pdf':
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted: raise HTTPException(422, 'Unlock the PDF before uploading.')
            if len(reader.pages) > 100: raise HTTPException(413, 'Use a PDF of at most 100 pages.')
            chunks = []
            for i, page in enumerate(reader.pages, 1):
                t = page.extract_text() or ''
                if t.strip(): chunks.append((f'page {i}',t))
                else: warnings.append(f'Page {i} has no extractable text. OCR is not included; its content was not read.')
        elif ext == '.docx':
            check_zip(data)
            doc = Document(io.BytesIO(data))
            chunks = []
            # iter_inner_content keeps paragraph/table ordering.
            for i, item in enumerate(doc.iter_inner_content(),1):
                if hasattr(item,'rows'):
                    t = '\n'.join(' | '.join(c.text for c in r.cells) for r in item.rows)
                else: t = item.text
                chunks.append((f'block {i}',t))
            warnings.append('DOCX extraction covers body paragraphs and tables, not embedded images, headers or text boxes.')
        else:
            check_zip(data)
            prs = Presentation(io.BytesIO(data))
            def texts(shapes):
                out=[]
                for sh in shapes:
                    if sh.has_text_frame: out.append(sh.text)
                    if sh.has_table: out.extend(' | '.join(c.text for c in row.cells) for row in sh.table.rows)
                    if hasattr(sh,'shapes'): out.extend(texts(sh.shapes))
                return out
            chunks=[(f'slide {i}','\n'.join(texts(s.shapes))) for i,s in enumerate(prs.slides,1)]
            warnings.append('PPTX extraction covers slide text/tables, not images, charts or speaker notes.')
    except HTTPException: raise
    except ValidationError:
        raise HTTPException(502, 'Media transcription was empty or malformed. Try a shorter clip with clear speech.') from None
    except Exception:
        raise HTTPException(422, f'Could not read {ext} file. Check its format; text files must use UTF-8.') from None
    chunks = [(loc,t.strip()) for loc,t in chunks if t.strip()]
    if not chunks: raise HTTPException(422, 'No readable text found. Scanned documents need OCR first.')
    return chunks,warnings,segments


def stamp(seconds):
    ms=round(seconds*1000)
    h,ms=divmod(ms,3600000); m,ms=divmod(ms,60000); s,ms=divmod(ms,1000)
    return f'{h:02}:{m:02}:{s:02},{ms:03}'


def srt(segments):
    return '\n\n'.join(f'{i}\n{stamp(s["start"])} --> {stamp(s["end"])}\n{s["text"]}'
        for i,s in enumerate(segments,1))+'\n'

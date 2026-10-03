import io
from typing import List, Literal, Optional
from pydantic import BaseModel
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.app.processing.image_utils import optimize_image_bytes


class PPTXOptimizationConfig(BaseModel):
    max_slides: Optional[int] = None
    include_images: bool = True
    max_image_dim: int = 1600
    image_quality: int = 80


class SlideText(BaseModel):
    slide_number: int
    text: str


class ExtractedImage(BaseModel):
    slide_number: int
    filename: str
    ext: str
    size_bytes: int
    data_b64: str


class OptimizedPPTXResult(BaseModel):
    doc_type: Literal["pptx"] = "pptx"
    filename: str
    original_size_bytes: int
    optimized_size_bytes: int
    slide_count: int
    slides_text: List[SlideText]
    images: List[ExtractedImage]


def optimize_pptx(
    pptx_bytes: bytes,
    filename: str,
    config: PPTXOptimizationConfig = PPTXOptimizationConfig(),
) -> OptimizedPPTXResult:
    prs = Presentation(io.BytesIO(pptx_bytes))
    total = len(prs.slides)

    slides_to_keep = (
        list(range(min(total, config.max_slides)))
        if config.max_slides
        else list(range(total))
    )

    slides_text: List[SlideText] = []
    images: List[ExtractedImage] = []
    total_text_bytes = 0

    for idx in slides_to_keep:
        slide = prs.slides[idx]
        slide_num = idx + 1
        texts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)

            if config.include_images and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = shape.image
                raw_blob = img.blob
                opt_bytes, final_ext, data_b64 = optimize_image_bytes(
                    image_bytes=raw_blob,
                    ext=img.ext,
                    max_dim=config.max_image_dim,
                    quality=config.image_quality,
                )
                images.append(
                    ExtractedImage(
                        slide_number=slide_num,
                        filename=f"slide{slide_num}_img{shape.shape_id}.{final_ext}",
                        ext=final_ext,
                        size_bytes=len(opt_bytes),
                        data_b64=data_b64,
                    )
                )

        joined_text = "\n".join(texts)
        total_text_bytes += len(joined_text.encode("utf-8"))
        slides_text.append(SlideText(slide_number=slide_num, text=joined_text))

    total_images_size = sum(img.size_bytes for img in images)
    optimized_size = total_text_bytes + total_images_size

    return OptimizedPPTXResult(
        filename=filename,
        original_size_bytes=len(pptx_bytes),
        optimized_size_bytes=optimized_size,
        slide_count=len(slides_to_keep),
        slides_text=slides_text,
        images=images,
    )
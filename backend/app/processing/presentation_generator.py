import io
import re
import urllib.request
from pathlib import Path
from typing import List, Dict, Any, Optional
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent

THEMES = {
    "amber": {
        "bg": RGBColor(17, 14, 8),
        "title": RGBColor(226, 162, 33),
        "subtitle": RGBColor(245, 183, 61),
        "body": RGBColor(223, 212, 192),
        "card": RGBColor(26, 21, 14),
        "muted": RGBColor(160, 150, 135),
    },
    "navy": {
        "bg": RGBColor(15, 23, 42),
        "title": RGBColor(56, 189, 248),
        "subtitle": RGBColor(125, 211, 252),
        "body": RGBColor(226, 232, 240),
        "card": RGBColor(30, 41, 59),
        "muted": RGBColor(148, 163, 184),
    },
    "light": {
        "bg": RGBColor(248, 250, 252),
        "title": RGBColor(15, 23, 42),
        "subtitle": RGBColor(2, 132, 199),
        "body": RGBColor(51, 65, 85),
        "card": RGBColor(255, 255, 255),
        "muted": RGBColor(100, 116, 139),
    },
}

IMG_REGEX = re.compile(
    r'(?:!\[(.*?)\]\((https?://[^\s\)]+)\))|'
    r'(?:\[(?:Image|Photo|Unsplash)[^\]]*\]\((https?://[^\s\)]+)\))|'
    r'(?:\*?Image(?:\s+Suggestion)?:\s*\[?(.*?)\]?\(?(https?://[^\s\)\*]+)\)?\*?)',
    re.IGNORECASE
)

def download_image_bytes(url: str, timeout: float = 3.0) -> Optional[bytes]:
    """Safely download image bytes with user-agent and timeout."""
    if not url or not url.startswith("http"):
        return None
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            if resp.status == 200:
                return resp.read()
    except Exception:
        pass
    return None

def resolve_logo_png_bytes(org_icon_url: str) -> Optional[bytes]:
    """Resolve organization logo to PNG bytes, converting SVG if necessary."""
    if not org_icon_url:
        org_icon_url = "/assets/logo.svg"

    try:
        # Check if it's a relative path to frontend assets
        if org_icon_url.startswith("/"):
            local_path = BASE_DIR / "frontend" / org_icon_url.lstrip("/")
            if not local_path.exists():
                local_path = BASE_DIR / "frontend" / "assets" / "logo.svg"
            if local_path.exists():
                if local_path.suffix.lower() == ".svg":
                    import pymupdf
                    doc = pymupdf.open(str(local_path))
                    return doc[0].get_pixmap().tobytes("png")
                else:
                    return local_path.read_bytes()
        elif org_icon_url.startswith("http"):
            raw_bytes = download_image_bytes(org_icon_url, timeout=3.0)
            if raw_bytes and org_icon_url.endswith(".svg"):
                import pymupdf
                doc = pymupdf.open(stream=raw_bytes, filetype="svg")
                return doc[0].get_pixmap().tobytes("png")
            return raw_bytes
    except Exception as e:
        print("Warning: Could not resolve logo for PPTX:", e)
    return None

def parse_slides_from_markdown(markdown_text: str) -> List[Dict[str, Any]]:
    """Parse raw markdown text into structured slide objects, ignoring subsequent deliverables."""
    clean_text = markdown_text.strip()
    if clean_text.startswith("```"):
        clean_text = re.sub(r"^```[a-zA-Z]*\n", "", clean_text)
        clean_text = re.sub(r"\n```$", "", clean_text)

    # Truncate before subsequent deliverables like Website or Mermaid Diagram
    deliv_match = re.search(
        r'\n#{1,3}\s+(?:\d+\.\s+)?(?:Deliverable:\s*)?(?:Website|Mermaid|Infographic|Diagram|Advisory|Executive)\b|\n```html|\n```mermaid',
        clean_text,
        re.IGNORECASE
    )
    if deliv_match:
        clean_text = clean_text[:deliv_match.start()].strip()

    # Strip top deliverable header if present (e.g. '## Deliverable: Presentation')
    clean_text = re.sub(
        r'^(?:#{1,3}\s+(?:\d+\.\s+)?(?:Deliverable:\s*)?(?:Presentation|Slides)\b[^\n]*\n+)',
        '',
        clean_text,
        flags=re.IGNORECASE
    ).strip()

    raw_slides = []
    if "\n---\n" in clean_text or "\n--- \n" in clean_text or "\n---\r\n" in clean_text:
        parts = re.split(r"\n\s*---\s*\n", clean_text)
        raw_slides = [p.strip() for p in parts if p.strip()]
    elif re.search(r"\n#{1,3}\s+Slide\s+\d+", clean_text, re.IGNORECASE):
        parts = re.split(r"\n(?=#{1,3}\s+Slide\s+\d+)", clean_text, flags=re.IGNORECASE)
        raw_slides = [p.strip() for p in parts if p.strip()]
    elif re.search(r"\n##\s+", clean_text):
        parts = re.split(r"\n(?=##\s+)", clean_text)
        raw_slides = [p.strip() for p in parts if p.strip()]
    else:
        paragraphs = [p.strip() for p in clean_text.split("\n\n") if p.strip()]
        if len(paragraphs) <= 3:
            raw_slides = [clean_text]
        else:
            raw_slides = [paragraphs[0]]
            for i in range(1, len(paragraphs), 2):
                chunk = "\n\n".join(paragraphs[i:i+2])
                raw_slides.append(chunk)

    slides = []
    for idx, slide_block in enumerate(raw_slides):
        # Skip blocks that are code blocks
        if slide_block.startswith("```html") or slide_block.startswith("```mermaid") or "<!DOCTYPE html>" in slide_block:
            continue

        lines = [l.strip() for l in slide_block.split("\n") if l.strip()]
        if not lines:
            continue

        slide_title = ""
        subtitle = ""
        bullets = []
        paragraphs = []
        image_url = ""
        image_caption = ""
        is_title_slide = (idx == 0) or ("title slide" in slide_block.lower()[:80])

        for line in lines:
            # Check for image links or suggestions
            m_img = IMG_REGEX.search(line)
            if m_img:
                groups = [g for g in m_img.groups() if g]
                for g in groups:
                    if g.startswith("http"):
                        image_url = g
                    elif not image_caption:
                        image_caption = g
                # Do NOT add this line as bullet
                continue

            # Ignore deliverable headers
            if re.match(r"^#{1,3}\s+(?:\d+\.\s+)?(?:Deliverable:\s*)?(?:Presentation|Slides)\b", line, flags=re.IGNORECASE):
                continue

            # Match titles
            if line.startswith("# ") or line.startswith("## ") or line.startswith("### "):
                title_cand = re.sub(r"^#{1,3}\s+", "", line).strip()
                title_cand = re.sub(r"^Slide\s+\d+:?\s*", "", title_cand, flags=re.IGNORECASE)
                if not slide_title:
                    slide_title = title_cand
                else:
                    bullets.append(title_cand)
            elif re.match(r"^\*\*(?:Slide\s+\d+:?\s*)?(.*?)\*\*$", line, re.IGNORECASE):
                m = re.match(r"^\*\*(?:Slide\s+\d+:?\s*)?(.*?)\*\*$", line, re.IGNORECASE)
                cand = m.group(1).strip()
                if not slide_title:
                    slide_title = cand
                else:
                    bullets.append(cand)
            elif line.startswith("* ") or line.startswith("- ") or line.startswith("• "):
                bullet_text = re.sub(r"^[\*\-•]\s+", "", line).strip()
                bullet_text = re.sub(r"^\*\*(.*?)\*\*:\s*", r"\1: ", bullet_text)
                bullets.append(bullet_text)
            elif re.match(r"^\d+\.\s+", line):
                bullet_text = re.sub(r"^\d+\.\s+", "", line).strip()
                bullets.append(bullet_text)
            else:
                if not slide_title:
                    slide_title = line.strip("#* ")
                elif is_title_slide and not subtitle:
                    subtitle = line.strip("#* ")
                else:
                    paragraphs.append(line)

        if not slide_title:
            slide_title = f"Slide {idx + 1}"

        slide_title = re.sub(r"\*\*|__", "", slide_title)

        slides.append({
            "slide_number": idx + 1,
            "title": slide_title,
            "subtitle": subtitle,
            "bullets": bullets,
            "paragraphs": paragraphs,
            "image_url": image_url,
            "image_caption": image_caption,
            "is_title_slide": is_title_slide,
        })

    return slides

def generate_pptx_bytes(
    slides: List[Dict[str, Any]],
    theme_name: str = "amber",
    org_name: str = "SatyaSetu",
    org_icon_url: str = "/assets/logo.svg"
) -> bytes:
    """Generate 16:9 widescreen PowerPoint presentation binary bytes with actual images and logo."""
    theme = THEMES.get(theme_name.lower(), THEMES["amber"])
    prs = pptx.Presentation()

    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Resolve logo PNG bytes once
    logo_bytes = resolve_logo_png_bytes(org_icon_url)

    for slide_data in slides:
        slide = prs.slides.add_slide(blank_layout)
        
        # Solid background
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = theme["bg"]

        is_title = slide_data.get("is_title_slide", False)

        if is_title:
            # Insert logo on Title Slide if available
            title_top = Inches(2.2)
            if logo_bytes:
                try:
                    slide.shapes.add_picture(
                        io.BytesIO(logo_bytes),
                        Inches(1.5),
                        Inches(1.0),
                        width=Inches(1.1),
                        height=Inches(1.1)
                    )
                    title_top = Inches(2.4)
                except Exception:
                    pass

            # Title Slide text frame
            title_box = slide.shapes.add_textbox(Inches(1.5), title_top, Inches(10.333), Inches(3.8))
            tf = title_box.text_frame
            tf.word_wrap = True

            p_title = tf.paragraphs[0]
            p_title.text = slide_data.get("title", "Presentation")
            p_title.font.bold = True
            p_title.font.size = Pt(44)
            p_title.font.color.rgb = theme["title"]
            p_title.space_after = Pt(16)

            subtitle_text = slide_data.get("subtitle", "")
            if not subtitle_text and slide_data.get("bullets"):
                subtitle_text = " | ".join(slide_data["bullets"][:2])

            if subtitle_text:
                p_sub = tf.add_paragraph()
                p_sub.text = subtitle_text
                p_sub.font.size = Pt(22)
                p_sub.font.color.rgb = theme["subtitle"]
                p_sub.space_after = Pt(24)

            p_org = tf.add_paragraph()
            display_org = org_name if org_name and org_name.lower() != "unincorporated" else "SatyaSetu AI"
            p_org.text = f"{display_org}  •  Presentation Deck"
            p_org.font.size = Pt(15)
            p_org.font.color.rgb = theme["muted"]

        else:
            # Content slide layout
            # Slide Header
            header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
            tf_head = header_box.text_frame
            tf_head.word_wrap = True
            p_head = tf_head.paragraphs[0]
            p_head.text = slide_data.get("title", f"Slide {slide_data.get('slide_number', '')}")
            p_head.font.bold = True
            p_head.font.size = Pt(32)
            p_head.font.color.rgb = theme["title"]

            # Check if this slide has an image
            image_url = slide_data.get("image_url", "")
            img_bytes = None
            if image_url:
                img_bytes = download_image_bytes(image_url)

            # Determine content box width based on image presence
            has_image = bool(img_bytes)
            content_width = Inches(6.8) if has_image else Inches(11.333)

            # Content Box (Left or Full width)
            content_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), content_width, Inches(4.5))
            tf_content = content_box.text_frame
            tf_content.word_wrap = True

            bullets = slide_data.get("bullets", [])
            paragraphs = slide_data.get("paragraphs", [])

            if bullets:
                for i, bullet in enumerate(bullets):
                    p = tf_content.paragraphs[0] if i == 0 else tf_content.add_paragraph()
                    cleaned_bullet = re.sub(r"\*\*|__", "", bullet)
                    p.text = f"•  {cleaned_bullet}"
                    p.font.size = Pt(19 if has_image else 20)
                    p.font.color.rgb = theme["body"]
                    p.space_after = Pt(12)
            elif paragraphs:
                for i, para in enumerate(paragraphs):
                    p = tf_content.paragraphs[0] if i == 0 else tf_content.add_paragraph()
                    p.text = re.sub(r"\*\*|__", "", para)
                    p.font.size = Pt(18)
                    p.font.color.rgb = theme["body"]
                    p.space_after = Pt(12)
            else:
                p = tf_content.paragraphs[0]
                p.text = "Key insights and discussion points."
                p.font.size = Pt(20)
                p.font.color.rgb = theme["muted"]

            # If image downloaded successfully, insert picture on the right side
            if has_image and img_bytes:
                try:
                    slide.shapes.add_picture(
                        io.BytesIO(img_bytes),
                        Inches(8.2),
                        Inches(2.2),
                        width=Inches(4.2),
                        height=Inches(4.0)
                    )
                except Exception:
                    pass

            # Insert logo in footer if available
            footer_left = Inches(1.0)
            if logo_bytes:
                try:
                    slide.shapes.add_picture(
                        io.BytesIO(logo_bytes),
                        Inches(1.0),
                        Inches(6.6),
                        height=Inches(0.4)
                    )
                    footer_left = Inches(1.6)
                except Exception:
                    pass

            # Slide Footer text
            footer_box = slide.shapes.add_textbox(footer_left, Inches(6.65), Inches(9.0), Inches(0.4))
            tf_footer = footer_box.text_frame
            p_foot = tf_footer.paragraphs[0]
            display_org = org_name if org_name and org_name.lower() != "unincorporated" else "SatyaSetu"
            p_foot.text = f"{display_org}  |  Slide {slide_data.get('slide_number', 1)}"
            p_foot.font.size = Pt(11)
            p_foot.font.color.rgb = theme["muted"]

    out_stream = io.BytesIO()
    prs.save(out_stream)
    return out_stream.getvalue()

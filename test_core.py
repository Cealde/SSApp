import io,json,os,tempfile,asyncio
from unittest.mock import AsyncMock,patch
os.environ['SATYA_DATA_DIR']=tempfile.mkdtemp(prefix='satya-test-')
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from app.main import app,provider
from app.schemas import FactSheet,Generated
from app.checks import validate_facts,validate_output,numbers
from app.ingest import extract,srt
from app.render import webpage
import httpx
client=TestClient(app)
SOURCE='The grant is INR 5,000. Selection is not guaranteed.'
SHEET={'facts':[{'id':'F1','statement':SOURCE,'evidence':[{'source_id':'S1','quote':SOURCE}]}],'conflicts':[]}
OUT={'title':'Grant announcement','blocks':[{'text':SOURCE,'fact_ids':['F1']}],'mermaid':''}
def prepared():
 r=client.post('/api/ingest',data={'text':SOURCE});assert r.status_code==200,r.text
 sid=r.json()['id']
 with patch.object(provider,'json',AsyncMock(return_value=SHEET)):
  r=client.post(f'/api/sessions/{sid}/facts');assert r.status_code==200,r.text
 assert client.put(f'/api/sessions/{sid}/facts',json={**SHEET,'reviewed':True}).status_code==200
 return sid

def test_round_trip_cache_approval_stale_facts():
 sid=prepared()
 with patch.object(provider,'json',AsyncMock(side_effect=[OUT,{'issues':[]}])) as mock:
  r=client.post(f'/api/sessions/{sid}/generate',json={'kind':'webpage'});assert r.status_code==200,r.text
  o=r.json();assert o['status']=='needs_human_review'
  assert client.post(f'/api/sessions/{sid}/generate',json={'kind':'webpage'}).json()['cached']
  assert mock.await_count==2
 assert client.post(f'/api/outputs/{o["id"]}/approve').json()['status']=='approved'
 assert 'INR 5,000' in client.get(f'/api/outputs/{o["id"]}/download').text
 changed=json.loads(json.dumps(SHEET));changed['facts'][0]['statement']='The grant amount is INR 5,000.'
 assert client.put(f'/api/sessions/{sid}/facts',json={**changed,'reviewed':True}).status_code==200
 assert client.post(f'/api/outputs/{o["id"]}/approve').status_code==409

def test_wrong_number_blocks():
 sid=prepared();bad=json.loads(json.dumps(OUT));bad['blocks'][0]['text']='Grant: INR 9,000.'
 with patch.object(provider,'json',AsyncMock(side_effect=[bad,{'issues':[]}])):
  o=client.post(f'/api/sessions/{sid}/generate',json={'kind':'summary'}).json()
 assert o['status']=='blocked'
 assert client.post(f'/api/outputs/{o["id"]}/approve').status_code==409

def test_semantic_negation_blocks():
 sid=prepared();bad=json.loads(json.dumps(OUT));bad['blocks'][0]['text']='Selection is guaranteed.'
 with patch.object(provider,'json',AsyncMock(side_effect=[bad,{'issues':[{'location':'block 1','reason':'Negation reversed.'}]}])):
  o=client.post(f'/api/sessions/{sid}/generate',json={'kind':'summary'}).json()
 assert o['status']=='blocked'

def test_invalid_evidence_unknown_citations():
 sheet=json.loads(json.dumps(SHEET));sheet['facts'][0]['evidence'][0]['quote']='invented'
 assert validate_facts(FactSheet.model_validate(sheet),[{'id':'S1','text':SOURCE}])
 out=json.loads(json.dumps(OUT));out['blocks'][0]['fact_ids']=['F999']
 assert validate_output(Generated.model_validate(out),FactSheet.model_validate(SHEET),'summary')

def test_unreviewed_conflict_gate():
 sid=client.post('/api/ingest',data={'text':SOURCE}).json()['id']
 assert client.post(f'/api/sessions/{sid}/generate',json={'kind':'summary'}).status_code==409
 assert client.put(f'/api/sessions/{sid}/facts',json={**SHEET,'conflicts':['Conflict'],'reviewed':True}).status_code==409

def test_multiple_office_files():
 d=Document();d.add_paragraph('Word source');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='Amount';t.cell(0,1).text='5000'
 db=io.BytesIO();d.save(db)
 p=Presentation();sl=p.slides.add_slide(p.slide_layouts[1]);sl.shapes.title.text='Slide source';pb=io.BytesIO();p.save(pb)
 r=client.post('/api/ingest',files=[('files',('a.docx',db.getvalue())),('files',('b.pptx',pb.getvalue())),('files',('c.txt',b'Text source'))])
 assert r.status_code==200,r.text
 text=' '.join(s['text'] for s in r.json()['sources'])
 assert all(t in text for t in ['Word source','Slide source','Text source','5000'])

def test_limits_bad_files():
 assert client.post('/api/ingest',data={'text':'a'*60001}).status_code==413
 assert client.post('/api/ingest',files={'files':('x.exe',b'content')}).status_code==415
 assert client.post('/api/ingest',files={'files':('x.pdf',b'bad')}).status_code==422

def test_html_escaping():
 h=webpage({'title':'<script>alert(1)</script>','blocks':[{'text':'<img src=x onerror=alert(1)>','fact_ids':['F1']}]},
 {'palette':'ocean','title_font':'Georgia','body_font':'Arial','layout':'article'})
 assert '<script>' not in h and '<img ' not in h and '&lt;script&gt;' in h

def test_media_srt():
 fake=AsyncMock();fake.json.return_value={'segments':[{'start':0.1,'end':1.2,'text':'Hello.'}],'visual_description':''}
 chunks,warnings,segs=asyncio.run(extract('hello.wav',b'fixture',fake))
 assert '00:00:00,100 --> 00:00:01,200' in srt(segs) and warnings

def test_indic_numbers():
 assert numbers('₹५,०००')==numbers('₹5,000')

def test_error_redacts_secret():
 old=provider.key;provider.key='private-test-secret'
 try:
  with patch('httpx.AsyncClient.request',AsyncMock(return_value=httpx.Response(429,json={'error':'private-test-secret'}))):
   r=client.post('/api/check-connection')
  assert r.status_code==429 and 'private-test-secret' not in r.text
 finally: provider.key=old

def test_rest_contract():
 old=provider.key;provider.key='private-test-secret'
 response={'candidates':[{'finishReason':'STOP','content':{'parts':[{'text':'{"ok":true}'}]}}]}
 try:
  with patch('httpx.AsyncClient.request',AsyncMock(return_value=httpx.Response(200,json=response))) as mock:
   assert client.post('/api/check-connection').json()['ok']
   assert mock.call_args.kwargs['headers']['x-goog-api-key']=='private-test-secret'
   assert mock.call_args.kwargs['json']['generationConfig']['responseMimeType']=='application/json'
 finally: provider.key=old
